#!/usr/bin/env python3
"""Readable Google Sheets/Excel exchange; source IDs remain hidden, lossless strings."""
import argparse
import csv
import json
import gzip
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation

VISIBLE = [('current_text', '현재 문구'), ('proposed_text', '수정 문안'), ('category', '분류'), ('context', '사용 위치'), ('notes', '참고'), ('action', '작업')]
HIDDEN = ['id', 'scope', 'source_path', 'source_line', 'placeholders', 'protected_markup', 'base_text_sha256', 'base_file_sha256']
FIELDS = [x[0] for x in VISIBLE] + HIDDEN
HEADERS = [x[1] for x in VISIBLE] + HIDDEN
SCOPES = [('game', '게임 문구'), ('comparison', '비교판 안내'), ('sound-lab', '소리 실험실'), ('dormant-content', '미노출 설정과 대사')]


def export(manifest, output):
    wb = Workbook()
    wb.remove(wb.active)
    for scope, title in SCOPES:
        ws = wb.create_sheet(title)
        ws.append(HEADERS)
        entries = [r for r in manifest['rows'] if r['scope'] == scope]
        for row in entries:
            ws.append([str(row.get(field, '')) for field in FIELDS])
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.zoomScale = 80
        for cell in ws[1]:
            cell.font = Font(color='FFFFFF', bold=True)
            cell.fill = PatternFill('solid', fgColor='243F3A')
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        ws.row_dimensions[1].height = 32
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.data_type = 's'  # Never interpret source/user text as a spreadsheet formula.
                cell.alignment = Alignment(wrap_text=True, vertical='top')
                cell.font = Font(size=11, color='172521')
            row[1].fill = PatternFill('solid', fgColor='FFF2CC')
            row[5].fill = PatternFill('solid', fgColor='FFF2CC')
            ws.row_dimensions[row[0].row].height = min(125, max(42, 16 * (len(str(row[0].value)) // 38 + 1)))
        for col, width in [('A', 48), ('B', 48), ('C', 22), ('D', 38), ('E', 45), ('F', 14)]:
            ws.column_dimensions[col].width = width
        ws.column_dimensions.group('G', 'N', hidden=True)
        ws['B1'].comment = Comment('노란 수정 문안만 입력하세요. 빈칸은 변경 없음입니다. 문구를 지우려면 작업 열을 clear로 지정하세요. {{expr:N}}와 HTML 태그는 유지하세요.', 'Text catalog')
        choices = DataValidation(type='list', formula1='"replace,clear"', allow_blank=True)
        choices.errorTitle = '작업 값'; choices.error = '빈칸, replace 또는 clear를 사용하세요.'; choices.showErrorMessage = True
        ws.add_data_validation(choices); choices.add(f'F2:F{max(2,ws.max_row)}')
    info = wb.create_sheet('사용 방법')
    instructions = [
        ('수정 방법', '게임 문구 탭의 노란 수정 문안 열에 바꿀 문장을 입력하세요. 빈칸은 변경하지 않겠다는 뜻입니다.'),
        ('구글 스프레드시트', 'Google Drive에 이 XLSX를 업로드하고 Google 스프레드시트로 열면 편집할 수 있습니다. 이 파일 자체는 Google 공유 문서가 아닙니다.'),
        ('전달 방법', '수정한 시트를 XLSX로 내려받아 전달하세요. 각 탭을 CSV로 내려받아도 됩니다. 동일 ID를 유지하세요.'),
        ('문구 삭제', '수정 문안을 비우고 작업 열을 clear로 지정하면 실제 빈 문자열로 바꿉니다.'),
        ('보호 문법', '{{expr:1}} 같은 자리표시자는 수·순서를 유지하세요. <li> 같은 기존 HTML 태그도 그대로 두세요. 따옴표와 줄바꿈은 쓸 수 있습니다.'),
        ('분리된 문장', '이어 붙인 문자열과 HTML의 강조 부분은 여러 행일 수 있습니다. 사용 위치와 참고를 확인하세요.'),
        ('원본 유지', '현재 문구·숨겨진 ID·해시 열은 수정하지 마세요. 수정 문안과 작업 열만 반영합니다.'),
        ('적용 검증', '기본 실행은 변경 보고서와 패치만 만듭니다. 원본 충돌·중복 ID·자리표시자·태그 손상을 거절합니다. --apply를 명시해야 게임 소스를 씁니다.'),
        ('미노출 탭', '저장된 설정·무기 영문 이름·현재 화면에 나오지 않는 대사입니다. 수정해도 현재 게임 화면에 곧바로 보이지 않습니다.'),
        ('포함 범위', '현재 게임 메뉴·무기 카드·표시 자세·캐릭터 이름/대사·접근성 문구와 비교판/소리 실험실을 분리했습니다.'),
        ('제외 범위', '개발 문서/성능 디버그/오류 로그/물리 식별자/주석/동결된 public/wb,corr 빌드. 이미지·음성 속 글은 자동 전사하지 않았습니다.'),
        ('기준 커밋', manifest['source_commit']),
        ('추출 시각', manifest['created_at']),
        ('행 수', str(len(manifest['rows']))),
    ]
    for row in instructions: info.append(row)
    info.column_dimensions['A'].width = 20; info.column_dimensions['B'].width = 105
    for row in info:
        for cell in row: cell.alignment = Alignment(wrap_text=True, vertical='top')
        info.row_dimensions[row[0].row].height = 44
    wb.save(output)


def to_csv(source, output):
    wb = load_workbook(source, data_only=False)
    records = []
    for _, title in SCOPES:
        if title not in wb.sheetnames:
            raise ValueError(f'Missing worksheet: {title}')
        ws = wb[title]
        if [c.value for c in ws[1]] != HEADERS:
            raise ValueError(f'Worksheet columns changed: {title}')
        for row in ws.iter_rows(min_row=2):
            if all(c.value is None for c in row): continue
            if any(c.data_type == 'f' for c in row):
                raise ValueError(f'Formula cell in {title}:{row[0].row}; use plain text')
            records.append({field: '' if cell.value is None else str(cell.value) for field, cell in zip(FIELDS, row)})
    with open(output, 'w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=FIELDS); writer.writeheader(); writer.writerows(records)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('command', choices=['export', 'csv'])
    parser.add_argument('input'); parser.add_argument('output')
    args = parser.parse_args()
    if args.command == 'export':
        data = Path(args.input).read_bytes()
        if args.input.endswith('.gz'): data = gzip.decompress(data)
        export(json.loads(data), args.output)
    else: to_csv(args.input, args.output)
    print(args.output)
